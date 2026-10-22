"""Shahi 1's total electricity and gas/LPG consumption, read from the
standalone monthly Energy Consumption Report workbook
(backend/data/utility/*Energy Consumption*.xlsx) -- a facility-wide report
covering Zain-1 and Shahi-1 together. Not derived from the OEE workbooks
or any other module here -- entirely its own small file, read off utility
bills each month by hand. (Was a single-sheet CSV export in an earlier
version of this workbook; now a proper 2-sheet .xlsx -- "Electricity
Consumption-26" and "Gas Consumption-26" -- so this reads a workbook, not
a flattened CSV, even though the column layout within each table is the
same content as before.)

Electricity ("Electricity Consumption-26" sheet) covers 4 power sources
(Gas Generator, Diesel Generator, K-Electric, Solar). Shahi-1's own unit
total isn't separately metered -- confirmed with the person maintaining
the sheet: the sheet computes it as (facility-wide Total Units Including
Solar) minus (Zain-1's own directly-metered total, itself Zain-1's base
reading plus its share of a shared "Compressor" pool), and reports that
result directly as "Shahi-1 Including Compressor" -- this module just
reads that column.

Only a facility-wide TOTAL is available for Shahi-1, not broken down by
source (Gas/Diesel/KE/Solar) -- the sheet's Zain-1/Shahi-1 split is only
ever computed on the combined total across all 4 sources, never per
source individually, so there's no way to attribute units/cost to one
source for Shahi-1 specifically (confirmed with the person running this
dashboard -- deliberately out of scope for now, not an oversight). Each
month's "sources" list is therefore the FACILITY-WIDE (Zain-1 + Shahi-1
combined) breakdown per source -- shown for context in the drill-down
alongside Shahi-1's own % share of the combined total, not an attempt to
split each source between the two facilities.

Electricity cost per month is DERIVED, not a column the sheet provides
directly: Shahi-1's units x that month's blended "Average unit cost
Including Solar" -- the same per-unit rate the sheet itself uses to turn
the facility-wide total into a blended cost, applied to Shahi-1's own
share of that total.

Gas/LPG ("Gas Consumption-26" sheet) is a separate fuel-use report, NOT
the same thing as the Electricity sheet's "Gas Generator" row (gas burned
specifically to run the generator) -- confirmed by the two sheets sharing
an identical Bill Amount for that one row/column pair (e.g. January: Rs.
467,088 in both), proving "Power Generation" here IS that same generator
expense, just measured in m3 of gas instead of electricity units
generated. To avoid re-showing that same expense twice under two
different headings, this module reads only "Shahi -1 Production" (direct
production-floor gas use, genuinely separate from the generator) for Gas,
and the "Shahi" sub-column of "Total LPG" for LPG -- both unambiguously
Shahi-1-only columns, unlike "Power Generation" which this module
deliberately skips.

Also NOT read here: Wood Pellets -- the sheet only reports one combined
Shahi+Zain total for it, with no Shahi-only split the way LPG has (per
the person running this dashboard, explicitly deferred rather than an
oversight), and Zain-1/Zain-2 Production, Zain-2 Power Generation, and
Zain's own LPG split -- out of scope, this dashboard covers Shahi-1 only.
"""
from __future__ import annotations

import threading
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).resolve().parent.parent / "data" / "utility"

MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]

ELECTRICITY_SHEET_NAME = "Electricity Consumption-26"
GAS_LPG_SHEET_NAME = "Gas Consumption-26"

# Fixed column positions within each sheet's own data rows -- confirmed
# directly against each sheet's own two/three-row header; there's no
# repeatable per-column label to search by the way labor_cost.py/gas.py
# do, since several columns share a bare "Units"/"Bill Amount" sub-header
# under different group headers.
COL_MONTH = 0
COL_GAS_UNITS = 1
COL_GAS_BILL_RS = 2
COL_GAS_UNIT_COST = 3
COL_DIESEL_UNITS = 4
COL_DIESEL_BILL_RS = 5
COL_DIESEL_UNIT_COST = 6
COL_KE_UNITS = 7
COL_KE_BILL_RS = 8
COL_KE_UNIT_COST = 9
COL_SOLAR_UNITS = 10
COL_TOTAL_UNITS_INCL_SOLAR = 11
COL_TOTAL_AMOUNT_RS = 13
COL_AVG_UNIT_COST_INCL_SOLAR = 15
COL_SHAHI_UNITS_INCL_COMPRESSOR = 21
COL_SHAHI_SHARE_PCT = 22

# (display name, units col, bill col, unit-cost col) -- Solar has no
# bill/unit-cost columns in the source sheet (it's free/self-generated),
# so those are None.
SOURCE_COLUMNS = [
    ("Gas Generator", COL_GAS_UNITS, COL_GAS_BILL_RS, COL_GAS_UNIT_COST),
    ("Diesel Generator", COL_DIESEL_UNITS, COL_DIESEL_BILL_RS, COL_DIESEL_UNIT_COST),
    ("K-Electric", COL_KE_UNITS, COL_KE_BILL_RS, COL_KE_UNIT_COST),
    ("Solar", COL_SOLAR_UNITS, None, None),
]

# "Gas Consumption-26" sheet's own columns -- see module docstring for
# why only these two (of many more on the sheet) are read.
COL_GAS_LPG_MONTH = 0
COL_SHAHI_GAS_M3 = 1
COL_SHAHI_GAS_BILL_RS = 2
COL_SHAHI_GAS_RATE = 3
COL_SHAHI_LPG_KG = 17
COL_SHAHI_LPG_AMOUNT_RS = 18

# Guards every read of _workbook_cache/_electricity_cache/_gas_lpg_cache
# below, not just their writes. FastAPI runs sync route handlers in a
# thread pool, and the unit page's own frontend code fires /electricity
# and /gas-lpg (and /overview, which now also reads electricity data for
# its cost-per-unit cards) concurrently via Promise.all on every page
# load -- openpyxl's Workbook/Worksheet objects aren't safe for
# concurrent iteration from multiple threads, confirmed by reproducing
# it directly: concurrent cold-start requests to /electricity and
# /gas-lpg left /electricity's sheet iteration silently reading zero
# rows (not an exception -- just corrupted internal state), which then
# cached as "no data" and stayed that way for the rest of the process,
# turning into a 404 (and, once other code started depending on it for
# non-optional data, a 500) on every subsequent request. A lock around
# the whole load+parse path is the simple, correct fix -- this data is
# read once per process lifetime (then served from the plain dict/list
# caches below, which concurrent reads of are safe), so serializing it
# costs nothing in practice.
# Reentrant: _load_electricity_rows/_load_gas_lpg_rows acquire this and
# then call _load_workbook, which acquires it again on the same thread.
_load_lock = threading.RLock()

_workbook_cache: openpyxl.Workbook | None = None
_workbook_loaded = False

_electricity_cache: list[dict] | None = None
_gas_lpg_cache: list[dict] | None = None


def _num(v) -> float | None:
    """Coerces a cell value to float, or None for blank/non-numeric
    placeholders (e.g. "Closed", "-") that a handful of cells elsewhere
    on these sheets use -- never raises."""
    if isinstance(v, (int, float)):
        return float(v)
    return None


def _round2(v: float | None) -> float | None:
    """round(v, 2), passing None through -- reading directly from the
    .xlsx (rather than a CSV export truncated by each cell's own display
    formatting) surfaces the sheet's true floating-point precision (e.g.
    81.87344434706398), which is more noise than signal for a Rs./unit
    figure."""
    return None if v is None else round(v, 2)


def _load_workbook() -> openpyxl.Workbook | None:
    global _workbook_cache, _workbook_loaded
    with _load_lock:
        if not _workbook_loaded:
            _workbook_loaded = True
            paths = list(DATA_DIR.glob("*Energy Consumption*.xlsx"))
            if paths:
                try:
                    _workbook_cache = openpyxl.load_workbook(paths[0], data_only=True)
                except OSError:
                    _workbook_cache = None
        return _workbook_cache


def _load_electricity_rows() -> list[dict]:
    global _electricity_cache
    with _load_lock:
        if _electricity_cache is not None:
            return _electricity_cache

        wb = _load_workbook()
        if wb is None or ELECTRICITY_SHEET_NAME not in wb.sheetnames:
            _electricity_cache = []
            return _electricity_cache
        ws = wb[ELECTRICITY_SHEET_NAME]

        results = []
        for row in ws.iter_rows(values_only=True):
            if not row or row[COL_MONTH] not in MONTH_NAMES:
                continue
            shahi_units = _num(row[COL_SHAHI_UNITS_INCL_COMPRESSOR])
            avg_cost_incl_solar = _num(row[COL_AVG_UNIT_COST_INCL_SOLAR])
            if shahi_units is None or avg_cost_incl_solar is None:
                continue  # month not filled in yet on the source sheet

            month_idx = MONTH_NAMES.index(row[COL_MONTH])
            sources = [
                {
                    "name": name,
                    "units": _round2(_num(row[units_col])),
                    "billRs": _round2(_num(row[bill_col])) if bill_col is not None else None,
                    "unitCost": _round2(_num(row[cost_col])) if cost_col is not None else None,
                }
                for name, units_col, bill_col, cost_col in SOURCE_COLUMNS
            ]
            results.append({
                "month": f"2026-{month_idx + 1:02d}",
                "monthLabel": f"{MONTH_NAMES[month_idx]} 2026",
                "unitsConsumed": round(shahi_units, 2),
                "costRs": round(shahi_units * avg_cost_incl_solar, 2),
                "sharePct": _round2(_num(row[COL_SHAHI_SHARE_PCT])),
                "facilityTotalUnits": _round2(_num(row[COL_TOTAL_UNITS_INCL_SOLAR])),
                "facilityTotalCostRs": _round2(_num(row[COL_TOTAL_AMOUNT_RS])),
                "facilityAvgUnitCost": round(avg_cost_incl_solar, 2),
                "sources": sources,
            })
        _electricity_cache = results
        return _electricity_cache


def compute_electricity_summary() -> dict | None:
    """Shahi-1's own monthly electricity units and (derived) cost -- see
    module docstring for how these are read/computed from the
    facility-wide sheet. Returns None if the source workbook isn't
    found."""
    rows = _load_electricity_rows()
    if not rows:
        return None
    return {"months": rows}


def _load_gas_lpg_rows() -> list[dict]:
    global _gas_lpg_cache
    with _load_lock:
        if _gas_lpg_cache is not None:
            return _gas_lpg_cache

        wb = _load_workbook()
        if wb is None or GAS_LPG_SHEET_NAME not in wb.sheetnames:
            _gas_lpg_cache = []
            return _gas_lpg_cache
        ws = wb[GAS_LPG_SHEET_NAME]

        results = []
        for row in ws.iter_rows(values_only=True):
            if not row or row[COL_GAS_LPG_MONTH] not in MONTH_NAMES:
                continue
            gas_m3 = _num(row[COL_SHAHI_GAS_M3])
            gas_bill_rs = _num(row[COL_SHAHI_GAS_BILL_RS])
            lpg_kg = _num(row[COL_SHAHI_LPG_KG])
            lpg_amount_rs = _num(row[COL_SHAHI_LPG_AMOUNT_RS])
            if gas_m3 is None and lpg_kg is None:
                continue  # month not filled in yet on the source sheet

            month_idx = MONTH_NAMES.index(row[COL_GAS_LPG_MONTH])
            gas_bill_rs = gas_bill_rs or 0.0
            lpg_amount_rs = lpg_amount_rs or 0.0
            results.append({
                "month": f"2026-{month_idx + 1:02d}",
                "monthLabel": f"{MONTH_NAMES[month_idx]} 2026",
                "gasM3": gas_m3,
                "gasCostRs": round(gas_bill_rs, 2) if gas_m3 is not None else None,
                "gasRatePerM3": round(gas_bill_rs / gas_m3, 2) if gas_m3 else None,
                "lpgKg": lpg_kg,
                "lpgCostRs": round(lpg_amount_rs, 2) if lpg_kg is not None else None,
                "lpgRatePerKg": round(lpg_amount_rs / lpg_kg, 2) if lpg_kg else None,
                "totalCostRs": round(gas_bill_rs + lpg_amount_rs, 2),
            })
        _gas_lpg_cache = results
        return _gas_lpg_cache


def compute_gas_lpg_summary() -> dict | None:
    """Shahi-1's own monthly direct-production Gas (m3) and LPG (KG)
    consumption/cost -- see module docstring for why this is "Shahi -1
    Production" gas specifically, not the sheet's "Power Generation" gas
    (that's the Electricity sheet's Gas Generator expense, already
    counted there), and why Wood Pellets isn't included (no Shahi-only
    split available yet). Returns None if the source workbook or this
    sheet isn't found."""
    rows = _load_gas_lpg_rows()
    if not rows:
        return None
    return {"months": rows}
