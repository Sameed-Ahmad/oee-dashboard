"""Labor cost per unit of output (packets for Packing Dept, KG for
Production Dept), read from the standalone monthly Labor Cost Analysis
workbook (backend/data/*Labor Cost Analysis*.xlsx) -- entirely separate
from the day-by-day OEE workbooks parser.py handles, and at a much coarser
grain (one figure per cost-center line per MONTH, not per day/shift), so
this is deliberately its own module.

Also exposes compute_output_reconciliation, which reuses this same
workbook's Output table to compare it against this dashboard's own
OEE-tracked output -- see that function's docstring for why the two
numbers are expected to differ (they measure different points in the
pipeline) rather than being a data-quality bug to fix.

Only Shahi-1 is covered -- the facility this dashboard's 10 products
belong to (Zain-1/Shahi-III are separate facilities also tracked in the
same workbook, but out of scope here).

Reads the workbook's own "Shahi-1 Report" sheet, which already computes
Output/Man-Hours/Labor Cost per cost-center line per month via formulas --
reusing those cached values rather than re-deriving from the raw source
sheet's own more complex per-month column blocks.

The workbook's cost-center line names don't map 1:1 onto this dashboard's
10 products (confirmed with the person maintaining both): Nimco is
reported as two separate packing lines ("R-10" and "R-20 / R-30") that
must be summed together; Coated Peanut and Namak Para share ONE combined
"Peanut & Namakpar" production line with no way to separate them at the
source, so both products show the same blended figure (mathematically,
splitting a blended cost proportionally by each product's own output
yields the identical cost-per-unit for both anyway -- there's no
information in this data to make them differ); Extruder has no
cost-center line anywhere in the workbook, so it has no figure at all.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

DATA_DIR = Path(__file__).resolve().parent.parent / "data"

REPORT_SHEET_NAME = "Shahi-1 Report"

TABLE_HEADERS = {
    "output": "Output",
    "manHours": "Total Man-Hours (Hrs)",
    "laborCost": "Total Labor Cost (Rs.)",
}

MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]
MONTH_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
MONTH_HEADER_RE = re.compile(r"^([A-Za-z]{3})-(\d{4})$")

# Product slug -> source line name(s) in the workbook to sum together.
# See module docstring for why Nimco/Coated Peanut/Namak Para aren't
# simple 1:1 matches, and why Extruder is absent entirely.
PRODUCT_SOURCE_LINES: dict[str, list[str]] = {
    "fryo": ["Packing -Fry -O"],
    "ishida": ["Packing - Ishida"],
    "pops": ["Packing - Pops & Weato"],
    "nimco": ["Packing - R-10", "Packing - R-20 / R-30"],
    "hnc-1": ["Production - H&C - 1"],
    "hnc-3": ["Production - H&C - 3"],
    "kuiper": ["Production - Kuiper"],
    "coated-peanut": ["Production - Peanut & Namakpar"],
    "namak-para": ["Production - Peanut & Namakpar"],
}

SHARED_PRODUCTS = {"coated-peanut", "namak-para"}

_workbook_cache: object | None = None
_workbook_loaded = False


def _num(v) -> float:
    return v if isinstance(v, (int, float)) else 0.0


def _load_workbook():
    global _workbook_cache, _workbook_loaded
    if not _workbook_loaded:
        _workbook_loaded = True
        candidates = sorted(DATA_DIR.glob("*Labor Cost Analysis*.xlsx"))
        _workbook_cache = openpyxl.load_workbook(candidates[0], data_only=True) if candidates else None
    return _workbook_cache


def _find_table(ws, label: str) -> tuple[int, dict[str, int]] | None:
    """Finds a table's anchor row by its exact label in column A, and
    returns (header_row, {"Jan-2026": column, ...}) for every month column
    found one row below it."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == label:
            header_row = r + 1
            months = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(row=header_row, column=c).value
                if isinstance(v, str) and MONTH_HEADER_RE.match(v):
                    months[v] = c
            return header_row, months
    return None


def _department_total_row(ws, table_header_row: int, department: str) -> int | None:
    """Finds `department`'s own TOTAL row within a table (e.g. the
    "Packing" row itself, already the sum of every one of its cost-center
    lines) -- distinct from _department_lines, which finds its children."""
    for r in range(table_header_row + 1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == department:
            return r
    return None


def _department_lines(ws, table_header_row: int, department: str) -> dict[str, int]:
    """Within a table (given its header row), finds `department`'s own row
    and returns {sub-department name: row number} for every line under it,
    stopping at the next populated column-A cell (the next department)."""
    lines: dict[str, int] = {}
    dept_row = None
    for r in range(table_header_row + 1, ws.max_row + 1):
        col1 = ws.cell(row=r, column=1).value
        if dept_row is None:
            if col1 == department:
                dept_row = r
            continue
        if col1 not in (None, ""):
            break
        sub = ws.cell(row=r, column=2).value
        if sub is None:
            break
        lines[str(sub).strip()] = r
    return lines


def compute_labor_cost_summary(slug: str) -> dict | None:
    """Returns {"sourceLines": [...], "shared": bool, "months": [...]}, or
    None if this product has no labor-cost line registered at all -- no
    workbook found, or genuinely absent (Extruder)."""
    source_lines = PRODUCT_SOURCE_LINES.get(slug)
    if not source_lines:
        return None

    wb = _load_workbook()
    if wb is None or REPORT_SHEET_NAME not in wb.sheetnames:
        return None
    ws = wb[REPORT_SHEET_NAME]

    output_table = _find_table(ws, TABLE_HEADERS["output"])
    hours_table = _find_table(ws, TABLE_HEADERS["manHours"])
    cost_table = _find_table(ws, TABLE_HEADERS["laborCost"])
    if not (output_table and hours_table and cost_table):
        return None

    department = "Production" if source_lines[0].startswith("Production") else "Packing"

    output_header_row, output_months = output_table
    hours_header_row, hours_months = hours_table
    cost_header_row, cost_months = cost_table

    output_lines = _department_lines(ws, output_header_row, department)
    hours_lines = _department_lines(ws, hours_header_row, department)
    cost_lines = _department_lines(ws, cost_header_row, department)

    unit = ""
    for line in source_lines:
        if line in output_lines:
            unit = ws.cell(row=output_lines[line], column=3).value or ""
            break

    common_headers = set(output_months) & set(hours_months) & set(cost_months)

    def sort_key(header: str) -> tuple[int, int]:
        mon, year = header.split("-")
        return int(year), MONTH_ABBR.index(mon)

    months_result = []
    for header in sorted(common_headers, key=sort_key):
        total_output = sum(
            _num(ws.cell(row=output_lines[line], column=output_months[header]).value)
            for line in source_lines if line in output_lines
        )
        if total_output <= 0:
            continue  # month not filled in yet on the source workbook

        total_hours = sum(
            _num(ws.cell(row=hours_lines[line], column=hours_months[header]).value)
            for line in source_lines if line in hours_lines
        )
        total_cost = sum(
            _num(ws.cell(row=cost_lines[line], column=cost_months[header]).value)
            for line in source_lines if line in cost_lines
        )

        mon, year = header.split("-")
        month_idx = MONTH_ABBR.index(mon)
        months_result.append({
            "month": f"{year}-{month_idx + 1:02d}",
            "monthLabel": f"{MONTH_NAMES[month_idx]} {year}",
            "manHours": round(total_hours, 1),
            "laborCost": round(total_cost, 2),
            "output": round(total_output, 1),
            "outputUnit": unit,
            "costPerUnit": round(total_cost / total_output, 4),
            "manHoursPerUnit": round(total_hours / total_output, 6),
        })

    return {
        "sourceLines": source_lines,
        "shared": slug in SHARED_PRODUCTS,
        "months": months_result,
    }


def compute_output_reconciliation(slug: str, records: list[dict]) -> dict | None:
    """Compares this dashboard's own OEE-tracked monthly output ("Produced"
    -- summed from `records`, the same day records every other endpoint
    uses) against the labor-cost workbook's own Output Input figure for the
    same source line(s) ("Received").

    These are two independently-maintained numbers for two different
    points in the pipeline -- machine-counted output vs. the quantity
    actually received into sellable finished-goods stock -- not the same
    quantity reported twice, so a gap between them is a genuine yield
    figure to track, not a data error to reconcile away.

    Uses `actualCounter` (the gross machine count, before quality-based
    rejection) rather than `stockTransferred` (already net of quality
    loss) for "Produced" -- Received can only ever be a subset of what was
    gross-produced, so actualCounter is the correct upper bound to compare
    against; stockTransferred would already have quality loss baked in,
    double-counting that loss against Received's own gap.

    Returns None if this product has no source line registered at all (see
    PRODUCT_SOURCE_LINES) -- same scope as compute_labor_cost_summary."""
    source_lines = PRODUCT_SOURCE_LINES.get(slug)
    if not source_lines:
        return None

    wb = _load_workbook()
    if wb is None or REPORT_SHEET_NAME not in wb.sheetnames:
        return None
    ws = wb[REPORT_SHEET_NAME]

    output_table = _find_table(ws, TABLE_HEADERS["output"])
    if not output_table:
        return None
    output_header_row, output_months = output_table

    department = "Production" if source_lines[0].startswith("Production") else "Packing"
    output_lines = _department_lines(ws, output_header_row, department)

    unit = ""
    for line in source_lines:
        if line in output_lines:
            unit = ws.cell(row=output_lines[line], column=3).value or ""
            break

    def sort_key(header: str) -> tuple[int, int]:
        mon, year = header.split("-")
        return int(year), MONTH_ABBR.index(mon)

    produced_by_month: dict[str, float] = {}
    for r in records:
        month_key = r["date"][:7]
        produced_by_month[month_key] = produced_by_month.get(month_key, 0.0) + r["actualCounter"]

    months_result = []
    for header in sorted(output_months, key=sort_key):
        received = sum(
            _num(ws.cell(row=output_lines[line], column=output_months[header]).value)
            for line in source_lines if line in output_lines
        )
        if received <= 0:
            continue  # month not filled in yet on the source workbook

        mon, year = header.split("-")
        month_idx = MONTH_ABBR.index(mon)
        month_key = f"{year}-{month_idx + 1:02d}"
        produced = produced_by_month.get(month_key, 0.0)
        if produced <= 0:
            continue  # no OEE-tracked data for this month (yet) -- nothing to compare

        gap = round(produced - received, 1)

        months_result.append({
            "month": month_key,
            "monthLabel": f"{MONTH_NAMES[month_idx]} {year}",
            "outputUnit": unit,
            "produced": round(produced, 1),
            "received": round(received, 1),
            "gap": gap,
            "gapPct": round(gap / produced * 100, 2),
        })

    return {
        "sourceLines": source_lines,
        "shared": slug in SHARED_PRODUCTS,
        "months": months_result,
    }


def compute_department_labor_cost(department: str) -> dict | None:
    """Total Labor Cost (Rs.) for an entire department ("Packing" or
    "Production"), summed across every month filled in on the source
    workbook so far. Reads the workbook's own department TOTAL row
    directly -- already the sum of every cost-center line in that
    department, including lines with no 1:1 mapping to one of this
    dashboard's 10 products (e.g. "Packing - Printing") -- rather than
    summing our own per-product figures, which would double-count Coated
    Peanut/Namak Para's shared line and miss unmapped lines entirely.

    Returns None if the workbook, this table, or this department's row
    can't be found, or no month has been filled in yet."""
    wb = _load_workbook()
    if wb is None or REPORT_SHEET_NAME not in wb.sheetnames:
        return None
    ws = wb[REPORT_SHEET_NAME]

    cost_table = _find_table(ws, TABLE_HEADERS["laborCost"])
    if not cost_table:
        return None
    header_row, months = cost_table

    row = _department_total_row(ws, header_row, department)
    if row is None:
        return None

    def sort_key(header: str) -> tuple[int, int]:
        mon, year = header.split("-")
        return int(year), MONTH_ABBR.index(mon)

    total = 0.0
    months_included = []
    for header in sorted(months, key=sort_key):
        value = _num(ws.cell(row=row, column=months[header]).value)
        if value <= 0:
            continue
        total += value
        mon, year = header.split("-")
        months_included.append(f"{year}-{MONTH_ABBR.index(mon) + 1:02d}")

    if not months_included:
        return None

    return {
        "totalRs": round(total, 2),
        "firstMonth": months_included[0],
        "lastMonth": months_included[-1],
    }
